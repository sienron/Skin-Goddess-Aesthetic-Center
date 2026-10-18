from pathlib import Path
import re

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

SOURCE = Path(r"C:\Users\Ron\.codex\attachments\598c8320-7085-4896-85cd-b4c2145395fd\Pasted text.txt")
OUTPUT = Path("SG_Appointment_Inventory_Test_Cases.docx")


def shade(cell, color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), color)


def borders(table):
    tbl_pr = table._tbl.tblPr
    b = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        item = OxmlElement(f"w:{edge}")
        item.set(qn("w:val"), "single")
        item.set(qn("w:sz"), "4")
        item.set(qn("w:color"), "D9D9D9")
        b.append(item)
    tbl_pr.append(b)


def cell_margins(cell, top=75, start=80, bottom=75, end=80):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    mar = tc_pr.first_child_found_in("w:tcMar")
    if mar is None:
        mar = OxmlElement("w:tcMar")
        tc_pr.append(mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    node = OxmlElement("w:tblHeader")
    node.set(qn("w:val"), "true")
    tr_pr.append(node)


def format_cell(cell, text, bold=False, white=False, center=False, size=7.5):
    cell.text = ""
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    cell_margins(cell)
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    run = p.add_run(text)
    run.bold = bold
    run.font.name = "Aptos"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor(255, 255, 255) if white else RGBColor(0, 0, 0)


def parse_cases(text):
    sections = []
    current = None
    for line in text.splitlines():
        match = re.match(r"^## ([A-N])\. (.+)$", line.strip())
        if match:
            current = {"letter": match.group(1), "title": match.group(2), "rows": []}
            sections.append(current)
            continue
        if current and line.startswith("|") and not re.match(r"^\|\s*(ID|---)", line):
            cells = [x.strip() for x in line.strip().strip("|").split("|")]
            if len(cells) == 4 and re.match(r"^[A-N]\d{2}$", cells[0]):
                current["rows"].append(cells)
    return sections


def set_width(cell, inches):
    cell.width = Inches(inches)
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    tc_w.set(qn("w:w"), str(int(inches * 1440)))
    tc_w.set(qn("w:type"), "dxa")


def add_footer(section):
    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0)
    r = p.add_run("SG Appointment and Inventory System  |  Test Case Workbook")
    r.font.name = "Aptos"
    r.font.size = Pt(8)
    r.font.color.rgb = RGBColor(89, 89, 89)


def main():
    sections = parse_cases(SOURCE.read_text(encoding="utf-8"))
    doc = Document()
    section = doc.sections[0]
    # US Letter portrait, sized for ordinary office printing.
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.55)
    section.bottom_margin = Inches(0.55)
    section.left_margin = Inches(0.55)
    section.right_margin = Inches(0.55)
    add_footer(section)
    styles = doc.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"]._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    styles["Normal"].font.size = Pt(9)
    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(4)
    r = title.add_run("SG Appointment and Inventory System")
    r.font.name = "Aptos Display"
    r.font.size = Pt(22)
    r.font.color.rgb = RGBColor(0, 0, 0)
    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.paragraph_format.space_after = Pt(12)
    sr = sub.add_run("Test Case Workbook")
    sr.bold = True
    sr.font.name = "Aptos"
    sr.font.size = Pt(13)
    sr.font.color.rgb = RGBColor(0, 0, 0)
    for index, group in enumerate(sections):
        heading = doc.add_paragraph()
        heading.paragraph_format.space_before = Pt(8 if index else 3)
        heading.paragraph_format.space_after = Pt(4)
        hr = heading.add_run(f"{group['letter']}. {group['title']}")
        hr.bold = True
        hr.font.name = "Aptos"
        hr.font.size = Pt(12)
        hr.font.color.rgb = RGBColor(0, 0, 0)
        table = doc.add_table(rows=1, cols=5)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        headers = ["ID", "Test Case", "Expected Result", "Paper Ref.", "Actual Result"]
        widths = [0.42, 2.0, 2.05, 0.75, 2.18]
        for c, header, width in zip(table.rows[0].cells, headers, widths):
            set_width(c, width)
            format_cell(c, header, bold=True, white=True, center=True, size=7.2)
            shade(c, "1F4E78")
        set_repeat_table_header(table.rows[0])
        for row_i, item in enumerate(group["rows"]):
            cells = table.add_row().cells
            values = item + [""]
            for col_i, (cell, value, width) in enumerate(zip(cells, values, widths)):
                set_width(cell, width)
                format_cell(cell, value, center=col_i in (0, 3), size=7.25)
                if row_i % 2 == 1:
                    shade(cell, "F4F8FB")
        borders(table)
    doc.core_properties.title = "SG Appointment and Inventory System Test Case Workbook"
    doc.core_properties.subject = "Functional and non-functional test cases"
    doc.save(OUTPUT)
    print(f"Created {OUTPUT} with {sum(len(x['rows']) for x in sections)} test cases.")


if __name__ == "__main__":
    main()
